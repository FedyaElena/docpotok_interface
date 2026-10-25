#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Сборка data/docs.js из Excel-файлов с результатами разноски.

Как пользоваться
----------------
1) Демо-данные (временные, для первого запуска):
       python tools/build_docs.py --demo
   Создает data/source/doc1.xlsx ... doc12.xlsx и собирает data/docs.js.

2) Свои данные:
   - положите Excel-файлы в data/source/ под именами doc1.xlsx, doc2.xlsx ... doc12.xlsx
     (номер файла = номер документа, который прототип узнает по имени загружаемого счета);
   - запустите:  python tools/build_docs.py

Формат Excel-файла (см. образцы, которые создает --demo)
-------------------------------------------------------
Лист «Разноска» (обязательный). Первая строка – заголовки, дальше проводки:
    Дата | Счет Дт | Счет Кт | Поставщик | Валюта | Сумма | Перевозка | Плательщик |
    Номер документа | Описание расхода | База распределения
Лист «Счет» (необязательный). Две колонки «Поле | Значение». Поддерживаются поля:
    Имя файла, Поставщик, ИНН, Номер счета, Дата, Сумма без НДС, Ставка НДС, Сумма НДС,
    Итого, Валюта, Тип объектов, Точность, Время обработки, Низкая уверенность
Если листа нет или поля не заполнены, значения считаются из листа «Разноска».
«Имя файла» – если исходный счет называется не doc7.pdf, а как-то иначе,
впишите его имя (можно несколько через «;»), и прототип узнает именно его.

Нужен пакет openpyxl:  pip install openpyxl
"""
import argparse
import json
import random
import re
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "data" / "source"
OUT = ROOT / "data" / "docs.js"

COLS = ["Дата", "Счет Дт", "Счет Кт", "Поставщик", "Валюта", "Сумма", "Перевозка",
        "Плательщик", "Номер документа", "Описание расхода", "База распределения"]
KEYS = ["date", "dt", "kt", "supplier", "currency", "amount", "transport", "payer",
        "docNo", "desc", "base"]


# ----------------------------------------------------------------------------
# Демо-данные
# ----------------------------------------------------------------------------
def fmt_date(d):
    return d

def demo_specs():
    """12 разных по сложности счетов."""
    P = "АО «ТехноСнаб»"
    return [
        dict(n=1, supplier="ООО «МорТранс»", inn="7703456789", number="СФ-2024-0847", date="18.04.2025",
             obj="контейнер", count=8, vat=20, payers=[P],
             fixed={"Фрахт": [36400, 41600, 39000, 33800, 44200, 39000, 36400, 41600],
                    "Хранение": [21900, 22450, 21700, 21500, 22050, 23000, 21850, 20870]},
             types=[("Фрахт", "по объему"), ("Хранение", "")], acc=98.7, sec=4.2),
        dict(n=2, supplier="ИП Сорокин А.В.", inn="502401234567", number="2025/04-112", date="17.04.2025",
             obj="контейнер", count=1, vat=0, payers=[P],
             types=[("Автоперевозка", "")], rng=(42000, 48000), acc=99.3, sec=2.6),
        dict(n=3, supplier="АО «Феникс Лог»", inn="7728123456", number="ФЛ-20250416", date="16.04.2025",
             obj="контейнер", count=12, vat=20, payers=[P, "ООО «Метиз-Импорт»"],
             types=[("Фрахт", "по объему"), ("Терминальная обработка", "поровну")], rng=(38000, 62000),
             acc=97.1, sec=8.9),
        dict(n=4, supplier="ООО «СкайКарго»", inn="7714998877", number="SKC-2025-891", date="15.04.2025",
             obj="отправка", count=5, vat=20, payers=[P],
             types=[("Воздушный фрахт", "по весу")], rng=(48000, 76000), acc=98.9, sec=3.8),
        dict(n=5, supplier="ТК «Даль»", inn="2536112233", number="ТКД-0419", date="14.04.2025",
             obj="контейнер", count=2, vat=20, payers=[P],
             types=[("Фрахт", "по объему"), ("Экспедирование", "")], rng=(12000, 31000), acc=99.1, sec=3.1),
        dict(n=6, supplier="ООО «ВостТранс»", inn="2721334455", number="ВТ-25-0441", date="14.04.2025",
             obj="контейнер", count=6, vat=20, payers=[P],
             types=[("Фрахт", "по объему"), ("Терминальная обработка", "поровну"), ("Демередж", "")],
             rng=(9000, 58000), acc=94.6, sec=6.4, low=["ИНН"]),
        dict(n=7, supplier="ООО «Балтик Шиппинг»", inn="7802556677", number="БШ-0312/25", date="12.04.2025",
             obj="контейнер", count=4, vat=20, payers=[P, "ООО «Метиз-Импорт»"],
             types=[("Фрахт", "по объему"), ("Таможенное оформление", "поровну"), ("Хранение", "")],
             rng=(14000, 64000), acc=97.8, sec=5.2),
        dict(n=8, supplier="ООО «ТрансКонтинент»", inn="7743990011", number="ТК-2025-0178", date="11.04.2025",
             obj="контейнер", count=10, vat=20, payers=[P],
             types=[("Фрахт", "по объему")], rng=(52000, 71000), acc=98.4, sec=5.7),
        dict(n=9, supplier="ООО «Порт-Сервис»", inn="2310445566", number="ПС-77/04", date="10.04.2025",
             obj="контейнер", count=3, vat=20, payers=[P],
             types=[("Хранение", ""), ("Демередж", ""), ("Погрузочно-разгрузочные работы", "по объему")],
             rng=(8000, 27000), acc=95.2, sec=4.9, low=["Дата"]),
        dict(n=10, supplier="ИП Мельников К.Д.", inn="165012345678", number="М-0409/1", date="09.04.2025",
             obj="машина", count=2, vat=0, payers=["ООО «Метиз-Импорт»"],
             types=[("Автодоставка", "по весу")], rng=(26000, 39000), acc=99.0, sec=2.9),
        dict(n=11, supplier="АО «СибЛогистик»", inn="5405778899", number="СЛ-25/1180", date="08.04.2025",
             obj="вагон", count=7, vat=20, payers=[P, "ООО «Метиз-Импорт»"],
             types=[("Фрахт", "по весу"), ("ЖД тариф", "по весу")], rng=(31000, 83000), acc=96.4, sec=7.5),
        dict(n=12, supplier="ООО «Альфа-Экспедиция»", inn="7725667788", number="АЭ-0405", date="05.04.2025",
             obj="контейнер", count=5, vat=20, payers=[P],
             types=[("Экспедирование", ""), ("Страхование груза", "по стоимости"), ("Фрахт", "по объему")],
             rng=(6000, 47000), acc=98.2, sec=5.0),
    ]


def make_demo_rows(s):
    rnd = random.Random(1000 + s["n"])
    rows = []
    for tname, base in s["types"]:
        for i in range(s["count"]):
            if "fixed" in s:
                amount = s["fixed"][tname][i]
            else:
                lo, hi = s["rng"]
                if tname in ("Хранение", "Демередж", "Страхование груза", "Экспедирование",
                             "Терминальная обработка", "Таможенное оформление"):
                    lo, hi = int(lo * 0.3), int(hi * 0.45)
                amount = int(round(rnd.uniform(lo, hi) / 50.0)) * 50
            payer = s["payers"][i % len(s["payers"])] if len(s["payers"]) > 1 else s["payers"][0]
            rows.append({
                "Дата": s["date"], "Счет Дт": "41", "Счет Кт": "60", "Поставщик": s["supplier"],
                "Валюта": "RUB", "Сумма": amount,
                "Перевозка": "ПВ-04%02d-%02d" % (s["n"] + 11, i + 1),
                "Плательщик": payer, "Номер документа": s["number"],
                "Описание расхода": tname, "База распределения": base,
            })
    return rows


def write_demo_sources():
    SRC.mkdir(parents=True, exist_ok=True)
    for s in demo_specs():
        rows = make_demo_rows(s)
        net = sum(r["Сумма"] for r in rows)
        vat = round(net * s["vat"] / 100)
        wb = Workbook()
        ws = wb.active
        ws.title = "Разноска"
        ws.append(COLS)
        for r in rows:
            ws.append([r[c] for c in COLS])
        widths = [12, 8, 8, 24, 8, 12, 16, 22, 18, 30, 20]
        for i, w in enumerate(widths):
            ws.column_dimensions[chr(65 + i)].width = w
        inv = wb.create_sheet("Счет")
        inv.append(["Поле", "Значение"])
        info = [
            ("Имя файла", "doc%d.pdf" % s["n"]),
            ("Поставщик", s["supplier"]), ("ИНН", s["inn"]), ("Номер счета", s["number"]),
            ("Дата", s["date"]), ("Сумма без НДС", net), ("Ставка НДС", "%d%%" % s["vat"]),
            ("Сумма НДС", vat), ("Итого", net + vat), ("Валюта", "RUB"),
            ("Тип объектов", s["obj"]), ("Точность", s["acc"]), ("Время обработки", s["sec"]),
            ("Низкая уверенность", ";".join(s.get("low", []))),
        ]
        for k, v in info:
            inv.append([k, v])
        inv.column_dimensions["A"].width = 22
        inv.column_dimensions["B"].width = 30
        wb.save(SRC / ("doc%d.xlsx" % s["n"]))


# ----------------------------------------------------------------------------
# Сборка docs.js из data/source/*.xlsx
# ----------------------------------------------------------------------------
def num(v, default=0):
    if v is None or v == "":
        return default
    if isinstance(v, (int, float)):
        return v
    t = re.sub(r"[^\d,.\-]", "", str(v)).replace(",", ".")
    try:
        return float(t) if "." in t else int(t)
    except ValueError:
        return default


def txt(v):
    if v is None:
        return ""
    if hasattr(v, "strftime"):
        return v.strftime("%d.%m.%Y")
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def read_doc(path):
    wb = load_workbook(path, data_only=True)
    ws = wb["Разноска"] if "Разноска" in wb.sheetnames else wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    header = [txt(c) for c in next(it)]
    idx = {}
    for c, k in zip(COLS, KEYS):
        if c not in header:
            sys.exit("%s: нет колонки «%s» на листе «Разноска»" % (path.name, c))
        idx[k] = header.index(c)
    rows = []
    for r in it:
        if r is None or all(v in (None, "") for v in r):
            continue
        row = {}
        for k in KEYS:
            v = r[idx[k]]
            row[k] = num(v) if k == "amount" else txt(v)
        rows.append(row)
    info = {}
    if "Счет" in wb.sheetnames:
        for r in wb["Счет"].iter_rows(min_row=2, values_only=True):
            if r and r[0]:
                info[txt(r[0])] = r[1]
    first = rows[0] if rows else {}
    net = num(info.get("Сумма без НДС"), sum(r["amount"] for r in rows))
    rate_raw = txt(info.get("Ставка НДС")) or "20%"
    rate = num(rate_raw.replace("%", ""), 20)
    vat = num(info.get("Сумма НДС"), round(net * rate / 100))
    total = num(info.get("Итого"), net + vat)
    objects = len({r["transport"] for r in rows if r["transport"]}) or len(rows)
    key = path.stem.lower()
    names = [n.strip() for n in txt(info.get("Имя файла")).split(";") if n.strip()]
    low = [n.strip() for n in txt(info.get("Низкая уверенность")).split(";") if n.strip()]
    return {
        "key": key,
        "fileNames": names,
        "invoice": {
            "supplier": txt(info.get("Поставщик")) or first.get("supplier", ""),
            "inn": txt(info.get("ИНН")),
            "number": txt(info.get("Номер счета")) or first.get("docNo", ""),
            "date": txt(info.get("Дата")) or first.get("date", ""),
            "net": net, "vatRate": rate, "vat": vat, "total": total,
            "currency": txt(info.get("Валюта")) or first.get("currency", "RUB"),
        },
        "objectLabel": txt(info.get("Тип объектов")) or "контейнер",
        "objects": objects,
        "accuracy": num(info.get("Точность"), 98.0),
        "seconds": num(info.get("Время обработки"), 4.0),
        "lowConfidence": low,
        "rows": rows,
    }


def natural(p):
    m = re.search(r"(\d+)", p.stem)
    return int(m.group(1)) if m else 0


def build():
    files = sorted(SRC.glob("*.xlsx"), key=natural)
    files = [f for f in files if not f.name.startswith("~$")]
    if not files:
        sys.exit("В data/source/ нет xlsx-файлов. Запустите с --demo или положите свои файлы.")
    docs = [read_doc(f) for f in files]
    OUT.write_text("// Файл создан tools/build_docs.py – не редактируйте вручную.\n"
                   "window.DOCS = " + json.dumps(docs, ensure_ascii=False, indent=1) + ";\n",
                   encoding="utf-8")
    print("Готово: %d документов -> %s" % (len(docs), OUT.relative_to(ROOT)))


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--demo", action="store_true", help="создать демо-данные в data/source и собрать docs.js")
    a = ap.parse_args()
    if a.demo:
        write_demo_sources()
    build()
